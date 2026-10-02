"""The communion liturgy in the Word file comes from liturgy_config.COMMUNION_BLOCKS
and prints exactly as before (slice 4 spec, Backend 6 `worship_service.py`;
Testing "Characterization first"; AC2). _legacy_communion is a verbatim copy
of worship_service._add_communion_liturgy as it was before slice 4a: the
characterization the refactor must keep, compared as XML in the same run, so
the python-docx version cannot matter."""
from io import BytesIO
from pathlib import Path

from docx import Document

import liturgy_config
import worship_service


def _legacy_communion(doc) -> None:
    """Add The Sacrament of the Lord's Supper liturgy (invitation, great thanksgiving, etc.)."""
    doc.add_paragraph("The Sacrament of the Lord's Supper", style="Heading 1")
    doc.add_paragraph()

    doc.add_paragraph("Invitation to the Table", style="Heading 2")
    doc.add_paragraph(
        "This is the table of our Lord Jesus Christ. It is not a reward for the righteous, "
        "but nourishment for those who hunger; not a prize for the strong, but grace for those who are weary. "
        "Here, blessing is not earned but received. All who seek to walk humbly with God, and who trust in God's mercy, "
        "are welcome at this table."
    )
    doc.add_paragraph()

    doc.add_paragraph("Great Thanksgiving", style="Heading 2")
    doc.add_paragraph("The Lord be with you.")
    p = doc.add_paragraph()
    p.add_run("And also with you.").bold = True
    doc.add_paragraph("Lift up your hearts.")
    p = doc.add_paragraph()
    p.add_run("We lift them up to the Lord.").bold = True
    doc.add_paragraph("Let us give thanks to the Lord our God.")
    p = doc.add_paragraph()
    p.add_run("It is right to give our thanks and praise.").bold = True
    doc.add_paragraph(
        "It is truly right and our greatest joy to give you thanks and praise, O God, "
        "creator of heaven and earth, for you have made us and all things, and in your love you hold us in life. "
        "And so we join the everlasting song:"
    )
    p = doc.add_paragraph()
    p.add_run(
        "Holy, holy, holy Lord, God of power and might, heaven and earth are full of your glory. "
        "Hosanna in the highest. Blessed is the one who comes in the name of the Lord. Hosanna in the highest."
    ).bold = True
    doc.add_paragraph(
        "You are holy, O God of majesty, and blessed is Jesus Christ, your Son, our Lord, "
        "who by his life, death, and resurrection has reconciled the world to you. On the night in which he gave himself up "
        "he took bread, gave thanks, broke it, and gave it to his disciples. And likewise the cup after supper. "
        "Remembering his death and resurrection, we offer ourselves in praise and thanksgiving. Therefore we proclaim the mystery of faith:"
    )
    doc.add_paragraph("Christ has died.")
    doc.add_paragraph("Christ is risen.")
    doc.add_paragraph("Christ will come again.")
    doc.add_paragraph()

    doc.add_paragraph("Words of Institution", style="Heading 2")
    doc.add_paragraph("[Words of institution as printed or as used.]")
    doc.add_paragraph()

    doc.add_paragraph("The Lord's Prayer", style="Heading 2")
    doc.add_paragraph("[The Lord's Prayer as printed.]")
    doc.add_paragraph()

    doc.add_paragraph("Breaking of the Bread and Communion", style="Heading 2")
    doc.add_paragraph(
        "The bread that we break is a sharing in the body of Christ. "
        "The cup that we bless is a sharing in the blood of Christ. Come, for all is ready."
    )
    doc.add_paragraph()

    doc.add_paragraph("Prayer After Communion", style="Heading 2")
    doc.add_paragraph(
        "Gracious God, we give you thanks that you have fed us at this table of grace, "
        "strengthening us not to win our lives, but to live them faithfully. Send us out to do justice, "
        "to love kindness, and to walk humbly with you, bearing your blessing into a world still hungry for hope, "
        "through Jesus Christ our Lord. Amen."
    )
    doc.add_paragraph()


def _service(**overrides):
    kwargs = dict(
        occasion="World Communion Sunday", date_display="October 04, 2026",
        hymns_by_slot={"opening": {"title": "Be Thou My Vision", "number": 450},
                       "response": {"title": "Come, Thou Fount", "number": 475}, "closing": None},
        liturgy={"prayers_of_the_people": "We pray.", "benediction": "Go in peace."},
        ot_ref="Isaiah 5:1-7", nt_ref="Matthew 21:33-46", include_communion=True)
    return BytesIO(worship_service.build_docx(**{**kwargs, **overrides}))      # bytes since slice 5a


def _paragraphs(buf) -> list[tuple[str, str, list[tuple[str, bool]]]]:
    return [(p.style.name, p.text, [(r.text, bool(r.bold)) for r in p.runs])
            for p in Document(buf).paragraphs]


def test_the_word_file_with_communion_is_unchanged(monkeypatch):
    built = _service().getvalue()
    current = Document(BytesIO(built)).element.body.xml
    paragraphs = _paragraphs(BytesIO(built))          # the new build's, taken before the swap
    monkeypatch.setattr(worship_service, "_add_communion_liturgy", _legacy_communion)
    legacy = Document(_service()).element.body.xml
    assert current == legacy
    assert ("Heading 1", "The Sacrament of the Lord's Supper",
            [("The Sacrament of the Lord's Supper", False)]) in paragraphs
    assert ("Normal", "And also with you.", [("And also with you.", True)]) in paragraphs


def test_communion_blocks_render_as_the_legacy_paragraphs():
    legacy, current = Document(), Document()
    _legacy_communion(legacy)
    worship_service._add_communion_liturgy(current)
    assert current.element.body.xml == legacy.element.body.xml
    styles = {"heading1": "Heading 1", "heading2": "Heading 2", "text": "Normal",
              "response": "Normal", "blank": "Normal"}
    assert [(p.style.name, p.text, any(r.bold for r in p.runs)) for p in current.paragraphs] == [
        (styles[b.style], b.text, b.style == "response") for b in liturgy_config.COMMUNION_BLOCKS]


def test_the_communion_text_lives_in_liturgy_config():
    source = Path(worship_service.__file__).read_text(encoding="utf-8")
    assert "Invitation to the Table" not in source and "Christ will come again." not in source
    assert "COMMUNION_BLOCKS" in source


def test_the_word_file_s_assurance_line_is_liturgy_config_s():
    paragraphs = _paragraphs(_service(liturgy={"assurance": "Leader: In Christ we are forgiven."}))
    at = next(i for i, (_style, text, _runs) in enumerate(paragraphs) if text == "Leader: In Christ we are forgiven.")
    assert paragraphs[at + 1] == ("Normal", liturgy_config.ASSURANCE_RESPONSE,
                                  [(liturgy_config.ASSURANCE_RESPONSE, True)])
    source = Path(worship_service.__file__).read_text(encoding="utf-8")
    assert "Thanks be to God" not in source

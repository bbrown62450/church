"""The Word file keeps Streamlit's layout (slice 5a spec, Testing
"Characterization comes first"; F §2.3.1; owner answer 3, 2026-10-01: exact
parity). legacy_build_docx is a verbatim copy of worship_service.build_docx as
it was before slice 5a (main at 934ffb9), with verbatim copies of the three
helpers it calls whose output the parity rests on (so a later change to them in
worship_service cannot move both sides at once; the communion helper, pinned by
test_communion_docx.py, is imported); each case renders both and compares the
document XML in the same run, so the python-docx version cannot matter. The only differences allowed are the
documented ones: the first reading's heading reads "First Reading", hymn
headings follow slots, a hymn without a number prints no "#None", the
readings come from resolve_readings (render_docx), a blank occasion prints no
second title line (owner decision A, 2026-10-02: Streamlit printed an empty
line break), and an Assurance with a typed "People:" label prints the text
before the label, then the fixed response once (owner decision B, 2026-10-02:
Streamlit printed the typed response inside the Leader line and then the fixed
one too)."""
import re
from datetime import date
from io import BytesIO
from typing import Any, Dict, List, Optional

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import liturgy_config as lc
import service_output as so
import worship_service
from liturgy_config import ASSURANCE_RESPONSE
from worship_service import _add_communion_liturgy

# --- worship_service's helpers as they were at 934ffb9, verbatim ---


def _add_leader_people_paragraph(doc, text: str) -> None:
    """Add paragraph(s): each Leader: (normal), each People: (bold). Supports multiple Leader/People pairs."""
    if not text or not text.strip():
        return
    # Find all "Leader:" and "People:" in order (case-insensitive)
    pattern = re.compile(r"\b(Leader|People):\s*", re.IGNORECASE)
    pos = 0
    parts = []
    for m in pattern.finditer(text):
        if m.start() > pos:
            parts.append(("", text[pos : m.start()].strip()))  # preamble if any
        role = "Leader" if m.group(1).lower() == "leader" else "People"
        end = pattern.search(text, m.end())
        content_end = end.start() if end else len(text)
        content = text[m.end() : content_end].strip()
        parts.append((role, content))
        pos = content_end
    if not parts:
        doc.add_paragraph(text)
        return
    for role, content in parts:
        if not content and role == "":
            continue
        if role == "People":
            p = doc.add_paragraph()
            p.add_run("People: ")
            r = p.add_run(content)
            r.bold = True
        else:
            line = ("Leader: " + content) if role == "Leader" else content
            if line:
                doc.add_paragraph(line)


def _add_assurance_paragraph(doc, leader_text: str) -> None:
    """Add Assurance: Leader line then liturgy_config.ASSURANCE_RESPONSE in bold
    (the one copy the 4b card shows too)."""
    leader_clean = (leader_text or "").strip()
    if leader_clean.startswith("Leader:"):
        leader_clean = leader_clean[7:].strip()
    if leader_clean:
        doc.add_paragraph("Leader: " + leader_clean)
    # Always add the congregational response
    p = doc.add_paragraph()
    r = p.add_run(ASSURANCE_RESPONSE)
    r.bold = True


def _add_custom_elements_after(
    doc,
    anchor: str,
    custom_elements: List[Dict[str, Any]],
) -> None:
    """Add any custom elements that are inserted after this anchor."""
    for ce in custom_elements:
        if ce.get("insert_after") == anchor and ce.get("label"):
            doc.add_paragraph(ce["label"], style="Heading 2")
            if ce.get("text"):
                doc.add_paragraph(ce["text"])
            doc.add_paragraph()


# --- worship_service.build_docx as it was at 934ffb9, verbatim ---


def legacy_build_docx(
    *,
    occasion: str,
    date: str,
    scriptures: List[str],
    hymns: List[Dict[str, str]],
    liturgy: Dict[str, str],
    include_placeholders: bool = True,
    sermon_title: Optional[str] = None,
    selected_ot_ref: Optional[str] = None,
    selected_nt_ref: Optional[str] = None,
    scripture_full_texts: Optional[Dict[str, str]] = None,
    include_sermon: bool = True,
    include_prayers_of_the_people: bool = True,
    include_communion: bool = False,
    custom_elements: Optional[List[Dict[str, Any]]] = None,
) -> BytesIO:
    custom = custom_elements or []
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Times New Roman"

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Worship Service\n{occasion}")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"
    if date:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(date).font.size = Pt(12)
    doc.add_paragraph()

    if liturgy.get("call_to_worship"):
        doc.add_paragraph("Call to Worship", style="Heading 2")
        _add_leader_people_paragraph(doc, liturgy["call_to_worship"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "call_to_worship", custom)

    if liturgy.get("opening_prayer"):
        doc.add_paragraph("Opening Prayer", style="Heading 2")
        doc.add_paragraph(liturgy["opening_prayer"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "opening_prayer", custom)

    if hymns:
        doc.add_paragraph("First Hymn", style="Heading 2")
        h = hymns[0]
        doc.add_paragraph(f"{h.get('title', '')} — #{h.get('number', '')}")
        doc.add_paragraph()
    _add_custom_elements_after(doc, "first_hymn", custom)

    if liturgy.get("prayer_of_confession"):
        doc.add_paragraph("Prayer of Confession", style="Heading 2")
        p = doc.add_paragraph()
        p.add_run(liturgy["prayer_of_confession"]).bold = True
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayer_of_confession", custom)

    if liturgy.get("assurance"):
        doc.add_paragraph("Assurance of Pardon", style="Heading 2")
        _add_assurance_paragraph(doc, liturgy["assurance"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "assurance", custom)

    if liturgy.get("prayer_for_illumination"):
        doc.add_paragraph("Prayer for Illumination", style="Heading 2")
        doc.add_paragraph(liturgy["prayer_for_illumination"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayer_for_illumination", custom)

    ot_ref = selected_ot_ref or (scriptures[0] if scriptures else None)
    if ot_ref:
        doc.add_paragraph("Old Testament Reading", style="Heading 2")
        doc.add_paragraph(ot_ref)
        doc.add_paragraph()
    _add_custom_elements_after(doc, "ot_reading", custom)

    nt_ref = selected_nt_ref or (scriptures[1] if len(scriptures) > 1 else None)
    if nt_ref:
        doc.add_paragraph("New Testament Reading", style="Heading 2")
        doc.add_paragraph(nt_ref)
        doc.add_paragraph()
    _add_custom_elements_after(doc, "nt_reading", custom)

    if include_sermon:
        doc.add_paragraph("Sermon Title", style="Heading 2")
        doc.add_paragraph(sermon_title.strip() if sermon_title and sermon_title.strip() else "[Sermon title]")
        doc.add_paragraph()
    _add_custom_elements_after(doc, "sermon", custom)

    doc.add_paragraph("Affirmation of Faith", style="Heading 2")
    doc.add_paragraph("Apostles' Creed")
    doc.add_paragraph()
    _add_custom_elements_after(doc, "affirmation_of_faith", custom)

    if len(hymns) > 1:
        doc.add_paragraph("Second Hymn", style="Heading 2")
        h = hymns[1]
        doc.add_paragraph(f"{h.get('title', '')} — #{h.get('number', '')}")
        doc.add_paragraph()
    _add_custom_elements_after(doc, "second_hymn", custom)

    if include_communion:
        _add_communion_liturgy(doc)
    _add_custom_elements_after(doc, "communion", custom)

    if include_prayers_of_the_people and liturgy.get("prayers_of_the_people"):
        doc.add_paragraph("Prayers of the People", style="Heading 2")
        doc.add_paragraph(liturgy["prayers_of_the_people"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayers_of_the_people", custom)

    if liturgy.get("offertory_prayer"):
        doc.add_paragraph("Offertory Prayer", style="Heading 2")
        doc.add_paragraph(liturgy["offertory_prayer"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "offertory_prayer", custom)

    if len(hymns) > 2:
        doc.add_paragraph("Third Hymn", style="Heading 2")
        h = hymns[2]
        doc.add_paragraph(f"{h.get('title', '')} — #{h.get('number', '')}")
        doc.add_paragraph()
    _add_custom_elements_after(doc, "third_hymn", custom)
    _add_custom_elements_after(doc, "benediction", custom)

    if liturgy.get("benediction"):
        doc.add_paragraph("Benediction", style="Heading 2")
        doc.add_paragraph(liturgy["benediction"])

    _add_custom_elements_after(doc, "end", custom)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


LITURGY = {
    "call_to_worship": "Leader: The Lord be with you. People: And also with you. Leader: Let us worship God.",
    "opening_prayer": "Gracious God, we gather in your name.",
    "prayer_of_confession": "Merciful God, we confess that we have not loved you.",
    "assurance": "Leader: In Jesus Christ we are forgiven.",
    "prayer_for_illumination": "Open our hearts by your Spirit.",
    "prayers_of_the_people": "We pray for the church and the world.",
    "offertory_prayer": "Receive these gifts.",
    "benediction": "Go in peace.",
}
HYMNS = [{"title": "Holy, Holy, Holy", "number": 138}, {"title": "Be Thou My Vision", "number": 450},
         {"title": "Amazing Grace", "number": 649}]
CUSTOM = [{"label": f"CE {key}", "text": f"Text after {key}.", "insert_after": key}
          for key, _label in lc.CUSTOM_PLACEMENTS]


def outline(content: bytes) -> list[tuple[str, str, bool]]:
    """(style, text, any run bold) for every paragraph."""
    return [(p.style.name, p.text, any(r.bold for r in p.runs)) for p in Document(BytesIO(content)).paragraphs]


def body_xml(content: bytes) -> str:
    return Document(BytesIO(content)).element.body.xml


def new_kwargs(*, hymns=HYMNS, ot="Isaiah 5:1-7", nt="Philippians 3:4b-14", **kw):
    slots = dict(zip(("opening", "response", "closing"), [*hymns, None, None, None]))
    return dict(occasion="World Communion Sunday", date_display="October 04, 2026",
                hymns_by_slot=slots, liturgy=dict(LITURGY), ot_ref=ot, nt_ref=nt,
                sermon_title="Living Water", include_communion=True, custom_elements=CUSTOM, **kw)


def legacy_kwargs(*, hymns=HYMNS, ot="Isaiah 5:1-7", nt="Philippians 3:4b-14", **kw):
    return dict(occasion="World Communion Sunday", date="October 04, 2026",
                scriptures=[s for s in (ot, nt) if s], hymns=list(hymns), liturgy=dict(LITURGY),
                selected_ot_ref=ot, selected_nt_ref=nt, sermon_title="Living Water",
                include_communion=True, custom_elements=CUSTOM, **kw)


@pytest.mark.parametrize("prayers", [False, True], ids=["bulletin", "pastor"])
def test_a_full_service_prints_exactly_as_streamlit_did_but_for_the_first_reading_heading(prayers):
    new = worship_service.build_docx(**new_kwargs(include_sermon=True, include_prayers_of_the_people=prayers))
    old = legacy_build_docx(**legacy_kwargs(include_sermon=True, include_prayers_of_the_people=prayers)).getvalue()
    assert isinstance(new, bytes)
    assert body_xml(new) == body_xml(old).replace("Old Testament Reading", "First Reading")
    assert Document(BytesIO(new)).styles.element.xml == Document(BytesIO(old)).styles.element.xml
    texts = [text for _style, text, _bold in outline(new)]
    assert "First Reading" in texts and "Old Testament Reading" not in texts
    assert ("Prayers of the People" in texts) is prayers


def test_a_sparse_service_prints_as_before_with_every_anchor():
    """No hymns, no readings, no communion, one section, an unknown liturgy key:
    every custom element still prints at its place (inv E5)."""
    liturgy = {"benediction": "Go in peace.", "notion_extra": "Never printed."}
    new = worship_service.build_docx(
        occasion="Ordinary Sunday", date_display="", hymns_by_slot={}, liturgy=liturgy, ot_ref=None, nt_ref=None,
        sermon_title="", include_sermon=True, include_prayers_of_the_people=True, include_communion=False,
        custom_elements=CUSTOM)
    old = legacy_build_docx(occasion="Ordinary Sunday", date="", scriptures=[], hymns=[], liturgy=liturgy,
                            sermon_title="", custom_elements=CUSTOM).getvalue()
    assert body_xml(new) == body_xml(old)
    texts = [text for _style, text, _bold in outline(new)]
    assert [t for t in texts if t.startswith("CE ")] == [f"CE {key}" for key, _ in lc.CUSTOM_PLACEMENTS]
    assert "[Sermon title]" in texts and "Never printed." not in texts


def test_the_title_date_fonts_and_bold_text_are_unchanged():
    doc = Document(BytesIO(worship_service.build_docx(**new_kwargs(include_sermon=True,
                                                                   include_prayers_of_the_people=True))))
    normal = doc.styles["Normal"].font
    assert (normal.name, normal.size) == ("Times New Roman", Pt(11))
    title, when = doc.paragraphs[0], doc.paragraphs[1]
    assert title.text == "Worship Service\nWorld Communion Sunday" and title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert (title.runs[0].bold, title.runs[0].font.size) == (True, Pt(16))
    assert when.text == "October 04, 2026" and when.runs[0].font.size == Pt(12)
    rows = [(p.style.name, p.text, any(r.bold for r in p.runs)) for p in doc.paragraphs]
    assert ("Normal", "People: And also with you.", True) in rows
    assert ("Normal", LITURGY["prayer_of_confession"], True) in rows
    assert ("Normal", lc.ASSURANCE_RESPONSE, True) in rows
    assert ("Heading 2", "Call to Worship", False) in rows


def test_hymn_headings_follow_slots_and_never_print_none():
    content = worship_service.build_docx(**new_kwargs(
        hymns=[None, {"title": "Be Thou My Vision", "number": 450}, {"title": "Old Favorite", "number": None}],
        include_sermon=True, include_prayers_of_the_people=False))
    texts = [text for _style, text, _bold in outline(content)]
    assert "First Hymn" not in texts                       # Streamlit printed the Response hymn here (inv D6)
    assert texts[texts.index("Second Hymn") + 1] == "Be Thou My Vision — #450"
    assert texts[texts.index("Third Hymn") + 1] == "Old Favorite"
    assert not any("#None" in t for t in texts)
    assert texts.index("CE first_hymn") < texts.index("Prayer of Confession")      # the anchor still prints


def _resolved(**kw):
    base = dict(service_date=date(2026, 10, 4), occasion="Nineteenth Sunday after Pentecost",
                scriptures=("Isaiah 5:1-7", "Psalm 80:7-15", "Philippians 3:4b-14", "Matthew 21:33-46"),
                hymns={"opening": so.ResolvedHymn("Holy, Holy, Holy", 138), "response": None, "closing": None},
                liturgy={"prayers_of_the_people": "We pray.", "benediction": "Go in peace."},
                sermon_title="", selected_ot_ref="", selected_nt_ref="", include_communion=False,
                custom_elements=(so.CustomElement("Anthem", "Choir", "sermon"),))
    return so.ResolvedService(**{**base, **kw})


def _after(texts: list[str], heading: str) -> str:
    return texts[texts.index(heading) + 1]


def test_render_docx_prints_the_readings_the_screen_shows():
    texts = [t for _s, t, _b in outline(so.render_docx(_resolved(selected_nt_ref="Matthew 21:33-40"), "pastor"))]
    assert _after(texts, "First Reading") == "Isaiah 5:1-7"
    assert _after(texts, "New Testament Reading") == "Philippians 3:4b-14"      # the stale pick is ignored
    easter = _resolved(scriptures=("Acts 10:34-43", "Psalm 118:1-2, 14-24", "Colossians 3:1-4", "John 20:1-18"))
    for variant in ("bulletin", "pastor"):
        texts = [t for _s, t, _b in outline(so.render_docx(easter, variant))]
        assert _after(texts, "First Reading") == "Acts 10:34-43"
        assert _after(texts, "New Testament Reading") == "Colossians 3:1-4"
        assert "Old Testament Reading" not in texts


def test_render_docx_variants():
    bulletin = [t for _s, t, _b in outline(so.render_docx(_resolved(), "bulletin"))]
    pastor = [t for _s, t, _b in outline(so.render_docx(_resolved(), "pastor"))]
    assert "Sermon Title" in bulletin and "[Sermon title]" in bulletin and "Prayers of the People" not in bulletin
    assert "Sermon Title" in pastor and _after(pastor, "Prayers of the People") == "We pray."
    assert _after(bulletin, "First Hymn") == "Holy, Holy, Holy — #138" and "Second Hymn" not in bulletin
    assert _after(bulletin, "Anthem") == "Choir"
    assert bulletin.index("Anthem") > bulletin.index("Sermon Title")
    communion = [t for _s, t, _b in outline(so.render_docx(_resolved(include_communion=True), "bulletin"))]
    assert communion.index(lc.COMMUNION_TITLE) > communion.index("Affirmation of Faith")


def test_without_python_docx_it_raises(monkeypatch):
    monkeypatch.setattr(worship_service, "Document", None)
    with pytest.raises(RuntimeError, match="python-docx is required"):
        so.render_docx(_resolved(), "bulletin")


def test_a_blank_occasion_prints_no_second_title_line_the_one_allowed_title_difference():
    """Owner decision A (2026-10-02): Streamlit's title run ended in an empty
    line break; the new one prints "Worship Service" alone. Everything else
    is as Streamlit printed it."""
    new = worship_service.build_docx(**{**new_kwargs(include_sermon=True, include_prayers_of_the_people=True),
                                        "occasion": ""})
    old = legacy_build_docx(**{**legacy_kwargs(include_sermon=True, include_prayers_of_the_people=True),
                               "occasion": ""}).getvalue()
    assert Document(BytesIO(new)).paragraphs[0].text == "Worship Service"
    assert Document(BytesIO(old)).paragraphs[0].text == "Worship Service\n"
    allowed = body_xml(old).replace("Old Testament Reading", "First Reading").replace(
        "<w:t>Worship Service</w:t>\n      <w:br/>", "<w:t>Worship Service</w:t>", 1)
    assert body_xml(new) == allowed


@pytest.mark.parametrize("typed, leader", [
    ("Leader: In Jesus Christ we are forgiven. People: Thanks be to God!", "In Jesus Christ we are forgiven."),
    ("Leader: In Jesus Christ we are forgiven.\npeople: Amen.", "In Jesus Christ we are forgiven."),
    ("In Christ we are forgiven.\nPEOPLE: Thanks be to God! Amen.", "In Christ we are forgiven."),
    ("People: Thanks be to God!", None),
])
def test_a_typed_people_label_in_the_assurance_prints_the_response_once(typed, leader):
    """Owner decision B (2026-10-02): the text prints up to the typed label
    (trimmed), then the fixed bold response once. Streamlit printed the typed
    response inside the Leader line; that is the one allowed difference."""
    liturgy = {"assurance": typed}
    kw = dict(date_display="", hymns_by_slot={}, ot_ref=None, nt_ref=None, include_sermon=False,
              include_prayers_of_the_people=False)
    rows = outline(worship_service.build_docx(occasion="Sunday", liturgy=liturgy, **kw))
    start = rows.index(("Heading 2", "Assurance of Pardon", False))
    expected = ([("Normal", f"Leader: {leader}", False)] if leader else []) + [
        ("Normal", lc.ASSURANCE_RESPONSE, True), ("Normal", "", False)]
    assert rows[start + 1:start + 1 + len(expected)] == expected
    assert [t for _s, t, _b in rows].count(lc.ASSURANCE_RESPONSE) == 1
    old = [t for _s, t, _b in outline(legacy_build_docx(occasion="Sunday", date="", scriptures=[], hymns=[],
                                                           liturgy=liturgy, include_sermon=False,
                                                           include_prayers_of_the_people=False).getvalue())]
    assert any("people:" in t.lower() and t != lc.ASSURANCE_RESPONSE for t in old)     # what Streamlit printed


def test_an_assurance_with_no_people_label_prints_as_streamlit_did():
    for typed in ("Leader: In Jesus Christ we are forgiven.", "God's peopled earth: forgiven.",
                  "Leader: The Lord's-People: forgiven."):
        liturgy = {"assurance": typed}
        new = worship_service.build_docx(occasion="Sunday", date_display="", hymns_by_slot={}, liturgy=liturgy,
                                         ot_ref=None, nt_ref=None, include_sermon=False,
                                         include_prayers_of_the_people=False)
        old = legacy_build_docx(occasion="Sunday", date="", scriptures=[], hymns=[], liturgy=liturgy,
                                include_sermon=False, include_prayers_of_the_people=False).getvalue()
        assert body_xml(new) == body_xml(old), typed

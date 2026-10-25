"""Text formatting (formatting spec, 2026-10-25; owner decisions 3 to 5):
the grammar (shared/formatting.json and the seeded shared/formatting_fuzz.json,
which the frontend's preview passes too), old text without a mark printing
exactly as before in every file and old text with one now formatting (the
owner's answer 3 of 2026-10-25), the Word copy, the printed bulletin (PDF and
Word) and the email's attachments reading the marks, nothing typed ever
becoming markup, the limits, the reviewer and the pews reading the words
without the marks, and Revise getting them."""
import dataclasses
import json
import random
import time
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.platypus import CondPageBreak

import affirmation as af
import formatting
import printed_bulletin as pb
import printed_docx
import printed_pdf
import service_bulletin as sb
import worship_service
from integrations import openai_client
from liturgy_config import ASSURANCE_RESPONSE, LIMITS
from printed_bulletin import Line, Span
from service_output import CustomElement
from tests.api_helpers import church_headers, make_api_client
from tests.test_build_docx_characterization import body_xml, legacy_build_docx, legacy_kwargs, new_kwargs
from tests.test_printed_bulletin import service
from usecases import archive, documents
from usecases import email as email_usecase

SHARED = Path(__file__).parent / "fixtures" / "shared"
FIXTURE = json.loads((SHARED / "formatting.json").read_text("utf-8"))
CASES = FIXTURE["cases"]
BOXES = ("prayer", "confession", "dialogue", "assurance", "announcement")

# Text saved before formatting, with characters a prayer may hold: none is a mark.
OLD = [
    "Psalm 23:1*", "#1 hymn: we sing it again", "5 * 3 * 2 = 30 blessings", "2*3*4", "A lone ** pair",
    "****", "a**b**c and **God**s love", "-dash first\n#hash first\n### three\n1.5 million\n1) one\n1234. four",
    "Line one\n\nLine three\n", "* not a list", "Back\\slash stays \\ here", "**open\nclose**", "# \n- ",
]


def _lines(box: str, text: str) -> list[Line]:
    if box == "prayer":
        return pb.formatted(text)
    if box == "confession":
        return pb.formatted(text, "bold", bold=True)
    if box == "dialogue":
        return pb._leader_people(text)
    if box == "assurance":
        return pb._assurance(text)
    return pb.formatted(text, "center")


def _as_fixture(line: Line) -> dict:
    """A line in the shared fixtures' form; "marker_bold" only on a bold marker (the Prayer of Confession's items)."""
    return {"style": line.style, "marker": line.marker, **({"marker_bold": True} if line.marker_bold else {}),
            "spans": [[s.text, ("b" if s.bold else "") + ("i" if s.italic else "")] for s in line.spans]}


# --- the grammar, shared with the frontend ---

def test_the_shared_fixture_covers_every_box_and_names_the_assurance_response():
    assert FIXTURE["assurance_response"] == ASSURANCE_RESPONSE
    assert {c["box"] for c in CASES} == set(BOXES)
    assert len(CASES) == 35 and len({c["name"] for c in CASES}) == 35


@pytest.mark.parametrize("case", CASES, ids=[c["name"] for c in CASES])
def test_the_printed_lines_are_the_shared_fixture_s(case):
    assert [_as_fixture(line) for line in _lines(case["box"], case["text"])] == case["lines"]


# The seeded random texts of shared/formatting_fuzz.json (plan review fixes, 2026-10-25): marks, spaces of every
# kind, labels and word characters thrown together, so the preview's copy of the rules meets what no one wrote a
# case for. The frontend's formatting.test.ts checks every case too.
_FUZZ_SEED, _FUZZ_CASES = 20261025, 300
_FUZZ_ATOMS = ("*", "**", "\\", "#", "# ", "## ", "- ", "-", "1. ", "123. ", "1234. ", ".", " ", "\t", "\u00a0", "\n",
               "\n", "a", "Z", "9", "_", "\u00e9", "People:", "Leader:", "people: ", "LEADER:", "word", "<b>", "&",
               "\u2003", "\u2028", "\u0085")


def fuzz_fixture() -> dict:
    """shared/formatting_fuzz.json as the printed bulletin prints it now."""
    rng = random.Random(_FUZZ_SEED)
    cases = []
    for _ in range(_FUZZ_CASES):
        box = rng.choice(BOXES)
        text = "".join(rng.choice(_FUZZ_ATOMS) for _ in range(rng.randint(1, 16)))
        cases.append({"box": box, "text": text, "lines": [_as_fixture(line) for line in _lines(box, text)]})
    return {"_about": "Seeded random texts (tests/test_formatting.py fuzz_fixture, seed 20261025) as the printed "
                      "bulletin prints them, in shared/formatting.json's form. Regenerate with "
                      "write_fuzz_fixture(); frontend/src/lib/formatting.test.ts checks every case.",
            "assurance_response": ASSURANCE_RESPONSE, "cases": cases}


def write_fuzz_fixture() -> None:
    """Rewrite shared/formatting_fuzz.json (one case a line, ASCII with escapes):
    `cd backend && ../.venv/bin/python -c "import tests.test_formatting as t; t.write_fuzz_fixture()"`."""
    data = fuzz_fixture()
    head = {k: v for k, v in data.items() if k != "cases"}
    rows = ",\n".join("  " + json.dumps(case) for case in data["cases"])
    (SHARED / "formatting_fuzz.json").write_text(json.dumps(head)[:-1] + ', "cases": [\n' + rows + "\n]}\n", "utf-8")


def test_the_seeded_random_texts_print_as_the_shared_fuzz_fixture_says():
    stored = json.loads((SHARED / "formatting_fuzz.json").read_text("utf-8"))
    assert stored == fuzz_fixture(), "the printed lines changed: rerun write_fuzz_fixture() and check the diff"


def test_the_marks():
    assert formatting.parse("**b** *i*") == [formatting.Block("text", (Span("b", True), Span(" "), Span("i", False, True)))]
    assert formatting.parse("# One\n## Two\n- three\n4. four") == [
        formatting.Block("heading1", (Span("One"),)), formatting.Block("heading2", (Span("Two"),)),
        formatting.Block("item", (Span("three"),), "•"), formatting.Block("item", (Span("four"),), "4.")]
    assert formatting.parse("# One\n- two", inline_only=True) == [formatting.Block("text", (Span("# One\n- two"),))]
    assert formatting.plain("# For all\n- **the** sick\n3. *the* lonely\n\\*") == "For all\n- the sick\n3. the lonely\n*"
    assert formatting.plain("Leader: **Come**.\n# x", inline_only=True) == "Leader: Come.\n# x"


@pytest.mark.parametrize("text", OLD)
def test_old_text_has_no_mark(text):
    assert not formatting.has_marks(text)
    assert formatting.plain(text) == text


def test_a_long_text_full_of_unclosed_marks_reads_quickly():
    for text in ("*a " * 6666, "**a " * 5000, "*" * LIMITS.max_section_text, "**x* " * 4000):
        started = time.monotonic()
        formatting.parse(text)
        assert time.monotonic() - started < 1.0


# --- old text without a mark prints exactly as before ---

def _legacy_formatted(text, style="body", *, bold=False):
    return [Line(style, (Span(text, bold=bold),))]


def _legacy_leader_people(text):
    """printed_bulletin._leader_people before formatting, verbatim."""
    import re
    matches = list(re.compile(r"\b(Leader|People):\s*", re.IGNORECASE).finditer(text))
    if not matches:
        return [Line("body", (Span(text),))]
    lines = []
    if matches[0].start() > 0 and text[:matches[0].start()].strip():
        lines.append(Line("body", (Span(text[:matches[0].start()].strip()),)))
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        content = text[m.end():end].strip()
        if m.group(1).lower() == "people":
            lines.append(Line("hanging", (Span("People: ", bold=True), Span(content, bold=True))))
        else:
            lines.append(Line("hanging", (Span("Leader: "), Span(content))))
    return lines


def _legacy_assurance(text):
    """printed_bulletin._assurance before formatting, verbatim."""
    leader = text.strip()
    people = pb._ASSURANCE_PEOPLE.search(leader)
    if people:
        leader = leader[: people.start()].strip()
    if leader.startswith("Leader:"):
        leader = leader[7:].strip()
    lines = [Line("hanging", (Span("Leader: "), Span(leader)))] if leader else []
    return [*lines, Line("bold", (Span(ASSURANCE_RESPONSE, bold=True),))]


def _old_service(text: str) -> pb.PrintedService:
    ps = service()
    liturgy = {key: text for key in ps.resolved.liturgy}
    liturgy["call_to_worship"] = f"Leader: {text}\nPeople: {text}"
    liturgy["assurance"] = f"Leader: {text}"
    resolved = dataclasses.replace(ps.resolved, liturgy=liturgy, custom_elements=(CustomElement("Anthem", text, "sermon"),),
                                   affirmation=af.Affirmation(True, "own", text))
    week = sb.ServiceBulletin(announcements=sb.Announcements(activities=text, prayer_concerns=text,
                                                             collection=text, other=text))
    return dataclasses.replace(ps, resolved=resolved, bulletin=week)


@pytest.mark.parametrize("text", OLD)
def test_old_text_prints_in_the_printed_bulletin_exactly_as_before(text, monkeypatch):
    ps = _old_service(text)
    lines = pb.order_of_worship(ps) + pb.announcements(ps)
    pdf, word = printed_pdf.render_pdf(ps), Document(BytesIO(printed_docx.render_docx(ps))).element.body.xml
    monkeypatch.setattr(pb, "formatted", _legacy_formatted)
    monkeypatch.setattr(pb, "_leader_people", _legacy_leader_people)
    monkeypatch.setattr(pb, "_assurance", _legacy_assurance)
    assert lines == pb.order_of_worship(ps) + pb.announcements(ps)
    assert pdf == printed_pdf.render_pdf(ps)                      # byte for byte (the PDF is invariant)
    assert word == Document(BytesIO(printed_docx.render_docx(ps))).element.body.xml


@pytest.mark.parametrize("text", OLD)
def test_old_text_prints_in_the_word_copy_exactly_as_streamlit_did(text):
    liturgy = {key: text for key in new_kwargs()["liturgy"]}
    liturgy["call_to_worship"] = f"Leader: {text}\nPeople: {text}"
    liturgy["assurance"] = f"Leader: {text}"
    custom = [{"label": "Anthem", "text": text, "insert_after": "sermon"}]
    new = new_kwargs(include_sermon=True, include_prayers_of_the_people=True)
    old = legacy_kwargs(include_sermon=True, include_prayers_of_the_people=True)
    for kwargs in (new, old):
        kwargs.update(liturgy=dict(liturgy), custom_elements=custom)
    assert body_xml(worship_service.build_docx(**new)) == body_xml(legacy_build_docx(**old).getvalue()).replace(
        "Old Testament Reading", "First Reading")


def test_old_own_affirmation_text_prints_in_the_word_copy_as_one_paragraph():
    doc = Document(BytesIO(worship_service.build_docx(**new_kwargs(affirmation=af.Affirmation(True, "own", OLD[7])))))
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("Affirmation of Faith") + 1] == OLD[7]


# --- changes to old text (the owner's answer 3 of 2026-10-25: accepted) ---
# Saved text with a mark of the grammar now prints formatted, and the Review step's preview shows it first. The
# owner's count (31 saved services): none with list-like lines in a prayer or a custom element, none in a
# bulletin, one with an asterisk before a letter.

CHANGED = [
    ("prayer", "- the sick\n- the lonely",
     [Line("item", (Span("the sick"),), marker="•"), Line("item", (Span("the lonely"),), marker="•")]),
    ("prayer", "1. We confess.\n2. We are forgiven.",
     [Line("item", (Span("We confess."),), marker="1."), Line("item", (Span("We are forgiven."),), marker="2.")]),
    ("prayer", "We pray for *all* people.", [Line("body", (Span("We pray for "), Span("all", italic=True), Span(" people.")))]),
    ("prayer", "Thanks be to God for **all** things.",
     [Line("body", (Span("Thanks be to God for "), Span("all", bold=True), Span(" things.")))]),
    ("announcement", "- Tuesday: Bible study\n- Wednesday: Choir",
     [Line("center", (Span("• Tuesday: Bible study"),)), Line("center", (Span("• Wednesday: Choir"),))]),
    ("prayer", "\\- not a list, \\*not italic\\*", [Line("body", (Span("- not a list, *not italic*"),))]),
]


@pytest.mark.parametrize(("box", "text", "lines"), CHANGED, ids=[text for _box, text, _lines in CHANGED])
def test_old_text_with_a_mark_now_prints_formatted(box, text, lines):
    assert formatting.has_marks(text)
    assert _lines(box, text) == lines


def test_an_old_dashed_list_now_prints_as_a_list_in_the_word_copy():
    liturgy = dict(new_kwargs()["liturgy"], prayers_of_the_people="For all who mourn:\n- the sick\n- the lonely")
    doc = Document(BytesIO(worship_service.build_docx(**{**new_kwargs(include_prayers_of_the_people=True),
                                                           "liturgy": liturgy})))
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("For all who mourn:"):][:3] == ["For all who mourn:", "•\tthe sick", "•\tthe lonely"]


# --- the Word copy ---

FORMATTED = "# For the world\nWe pray for **peace** and *justice*.\n## Near\n- the sick\n12. the lonely"


def _paragraph(doc, text):
    return next(p for p in doc.paragraphs if p.text == text)


def test_the_word_copy_reads_the_marks():
    liturgy = dict(new_kwargs()["liturgy"], opening_prayer=FORMATTED, prayer_of_confession="We confess *our sin*.",
                   call_to_worship="Leader: Come, **all**.\nPeople: We come *gladly*.")
    doc = Document(BytesIO(worship_service.build_docx(**{**new_kwargs(), "liturgy": liturgy})))
    heading = _paragraph(doc, "For the world")
    assert heading.style.name == "Normal" and heading.paragraph_format.keep_with_next
    assert [(r.bold, r.font.size) for r in heading.runs] == [(True, Pt(13))]
    assert [(r.bold, r.font.size) for r in _paragraph(doc, "Near").runs] == [(True, Pt(12))]
    text = _paragraph(doc, "We pray for peace and justice.")
    assert [(r.text, r.bold, r.italic) for r in text.runs] == [
        ("We pray for ", None, None), ("peace", True, None), (" and ", None, None), ("justice", None, True),
        (".", None, None)]
    for item, marker, hang in (("•\tthe sick", "•\t", 0.25), ("12.\tthe lonely", "12.\t", 0.375)):
        p = _paragraph(doc, item)
        assert p.runs[0].text == marker
        assert (p.paragraph_format.left_indent, p.paragraph_format.first_line_indent) == (Inches(hang), -Inches(hang))
    confession = _paragraph(doc, "We confess our sin.")
    assert [(r.text, r.bold, r.italic) for r in confession.runs] == [
        ("We confess ", True, None), ("our sin", True, True), (".", True, None)]
    people = _paragraph(doc, "People: We come gladly.")
    assert [(r.text, r.bold, r.italic) for r in people.runs] == [
        ("People: ", None, None), ("We come ", True, None), ("gladly", True, True), (".", True, None)]
    assert [(r.text, r.bold) for r in _paragraph(doc, "Leader: Come, all.").runs] == [
        ("Leader: Come, ", None), ("all", True), (".", None)]


# Every Word copy box that prints through worship_service._add_text, each with its own words (fix agent, finding 1).
_TEXT_BOXES = ("opening_prayer", "prayer_for_illumination", "prayers_of_the_people", "offertory_prayer", "benediction")
_ANCHORS = ("call_to_worship", "opening_prayer", "first_hymn", "prayer_of_confession", "assurance",
            "prayer_for_illumination", "ot_reading", "nt_reading", "sermon", "affirmation_of_faith", "second_hymn",
            "communion", "prayers_of_the_people", "offertory_prayer", "third_hymn", "benediction", "end")


def _boxed(name: str) -> str:
    return f"# {name} heading\nWe pray **{name}**.\n- {name} item"


def _assert_reads_the_marks(doc, name: str, *, bold: bool = False) -> None:
    heading = _paragraph(doc, f"{name} heading")
    assert heading.style.name == "Normal" and heading.paragraph_format.keep_with_next, name
    assert [(r.bold, r.font.size) for r in heading.runs] == [(True, Pt(13))], name
    assert [(r.text, r.bold) for r in _paragraph(doc, f"We pray {name}.").runs] == [
        ("We pray ", True if bold else None), (name, True), (".", True if bold else None)], name
    item = _paragraph(doc, f"•\t{name} item")
    assert [(r.text, r.bold) for r in item.runs] == [("•\t", True if bold else None), (f"{name} item",
                                                                                      True if bold else None)], name


def test_every_word_copy_box_reads_the_marks():
    liturgy = dict(new_kwargs()["liturgy"], **{key: _boxed(key) for key in _TEXT_BOXES},
                   prayer_of_confession=_boxed("prayer_of_confession"))
    custom = [{"label": f"Custom {anchor}", "text": _boxed(f"custom {anchor}"), "insert_after": anchor}
              for anchor in _ANCHORS]
    doc = Document(BytesIO(worship_service.build_docx(**{**new_kwargs(include_prayers_of_the_people=True),
                                                           "liturgy": liturgy, "custom_elements": custom})))
    for key in _TEXT_BOXES:
        _assert_reads_the_marks(doc, key)
    _assert_reads_the_marks(doc, "prayer_of_confession", bold=True)
    for anchor in _ANCHORS:
        _assert_reads_the_marks(doc, f"custom {anchor}")


def test_own_affirmation_text_reads_the_marks_plain_in_the_word_copy():
    """The spec's table: the Word copy prints own Affirmation text plain, as today (only the printed bulletin makes
    it all bold): its marks read, nothing else bold (finding 2)."""
    text = "I believe in **God** and *grace*.\n- the church\n# Amen"
    doc = Document(BytesIO(worship_service.build_docx(**new_kwargs(affirmation=af.Affirmation(True, "own", text)))))
    assert [(r.text, r.bold, r.italic) for r in _paragraph(doc, "I believe in God and grace.").runs] == [
        ("I believe in ", None, None), ("God", True, None), (" and ", None, None), ("grace", None, True),
        (".", None, None)]
    assert [(r.text, r.bold) for r in _paragraph(doc, "•\tthe church").runs] == [("•\t", None), ("the church", None)]
    assert [(r.bold, r.font.size) for r in _paragraph(doc, "Amen").runs] == [(True, Pt(13))]


# --- the printed bulletin and the email ---

@pytest.fixture
def church(make_user, make_church):
    return make_church(name="Grace", owner_user_id=make_user(email="pastor@example.com"))


def _input(**liturgy) -> archive.ServiceInput:
    return archive.ServiceInput(service_date=date(2026, 10, 4), liturgy=liturgy)


def _pdf_text(content: bytes) -> str:
    """The PDF's words; pypdf reads the standard fonts' bullet (WinAnsi code 0x7F) as "\x7f"."""
    pages = PdfReader(BytesIO(content)).pages
    return " ".join(" ".join(page.extract_text().replace("\x7f", "•").split()) for page in pages)


def test_the_printed_bulletin_reads_the_marks_in_pdf_and_word(church):
    pdf = _pdf_text(documents.build_printed(church, _input(opening_prayer=FORMATTED), "pdf", None,
                                            charge=lambda n: None).content)
    assert "For the world We pray for peace and justice. Near • the sick 12. the lonely" in pdf
    assert "**" not in pdf and "# " not in pdf
    word = Document(BytesIO(documents.build_printed(church, _input(opening_prayer=FORMATTED), "docx", None,
                                                    charge=lambda n: None).content))
    # The printed bulletin's box headings: 12 pt and 11 pt (the owner's answer 1 of 2026-10-25).
    assert [(r.bold, r.font.size) for r in _paragraph(word, "For the world").runs] == [(True, Pt(12))]
    assert [(r.bold, r.font.size) for r in _paragraph(word, "Near").runs] == [(True, Pt(11))]
    # A heading stays with what follows it (printed_docx._KEEP_WITH_NEXT), as the PDF's does (printed_pdf.KEEP).
    assert _paragraph(word, "For the world").paragraph_format.keep_with_next
    assert _paragraph(word, "Near").paragraph_format.keep_with_next
    assert not _paragraph(word, "We pray for peace and justice.").paragraph_format.keep_with_next
    for text, hang in (("•\tthe sick", 0.25), ("12.\tthe lonely", 0.375)):
        item = _paragraph(word, text)
        assert (item.paragraph_format.left_indent, item.paragraph_format.first_line_indent) == (Inches(hang),
                                                                                               -Inches(hang))
    assert [(r.text, r.bold) for r in _paragraph(word, "We pray for peace and justice.").runs][1] == ("peace", True)


def _pdf_runs(content: bytes) -> list[tuple[str, str]]:
    """The PDF's text pieces with their fonts, in drawing order."""
    runs = []
    for page in PdfReader(BytesIO(content)).pages:
        page.extract_text(visitor_text=lambda text, _cm, _tm, font, _size: runs.append(
            (text.strip().replace("\x7f", "•"), font.get("/BaseFont") if font else None)) if text.strip() else None)
    return runs


def test_a_box_all_bold_prints_its_list_markers_bold_in_pdf_and_word(church):
    """The spec's table: the Prayer of Confession is every run bold in all three files, its markers too (the Word copy
    bolds them: test_every_word_copy_box_reads_the_marks), and so is own Affirmation text in the printed bulletin;
    another box's markers stay plain (finding 4)."""
    typed = archive.ServiceInput(service_date=date(2026, 10, 4),
                                 liturgy=dict(prayer_of_confession="We confess:\n7. our pride",
                                              opening_prayer="We pray:\n8. the sick"),
                                 affirmation=af.Affirmation(True, "own", "We trust:\n- in God"))

    def build(fmt):
        return documents.build_printed(church, typed, fmt, None, charge=lambda n: None).content

    word = Document(BytesIO(build("docx")))
    assert [(r.text, r.bold) for r in _paragraph(word, "7.\tour pride").runs] == [("7.\t", True), ("our pride", True)]
    assert [(r.text, r.bold) for r in _paragraph(word, "•\tin God").runs] == [("•\t", True), ("in God", True)]
    assert [(r.text, r.bold) for r in _paragraph(word, "8.\tthe sick").runs] == [("8.\t", None), ("the sick", None)]
    runs = _pdf_runs(build("pdf"))
    for marker, words, font in (("7.", "our pride", "/Times-Bold"), ("•", "in God", "/Times-Bold"),
                                ("8.", "the sick", "/Times-Roman")):
        at = runs.index((words, font))
        assert runs[at - 1] == (marker, font)


_NUMBERED = "1. one\n10. ten\n100. hundred\n999. most"


def test_a_list_s_words_line_up_from_one_to_three_digit_numbers_in_the_pdf():
    """A numbered item hangs 27 pt (formatting.hanging_points): the widest number ("999.") and the PDF's bullet gap
    fit, so no number pushes its words right; a bulleted item keeps 18 pt (finding 6)."""
    for box in ("opening_prayer", "prayer_of_confession"):
        ps = service()
        ps = dataclasses.replace(ps, resolved=dataclasses.replace(
            ps.resolved, liturgy={**ps.resolved.liturgy, box: _NUMBERED + "\n- bullet"}))
        starts = {}
        for page in PdfReader(BytesIO(printed_pdf.render_pdf(ps))).pages:
            page.extract_text(visitor_text=lambda text, cm, tm, _font, _size: starts.setdefault(
                text.strip(), round(tm[4] + cm[4], 2)))
        words = [starts[w] - starts["1."] for w in ("one", "ten", "hundred", "most", "bullet")]
        assert words == [27.0, 27.0, 27.0, 27.0, 18.0], box


def test_a_numbered_item_hangs_wider_than_a_bulleted_one_in_both_word_files():
    """Word's hanging indent is the marker's tab stop: "100." (19.25 pt) passes 0.25 in, which would send its words
    to the next default stop, so a number hangs 0.375 in (finding 6)."""
    ps = service()
    ps = dataclasses.replace(ps, resolved=dataclasses.replace(ps.resolved, liturgy={**ps.resolved.liturgy,
                                                                                    "opening_prayer": _NUMBERED}))
    liturgy = dict(new_kwargs()["liturgy"], opening_prayer=_NUMBERED)
    for content in (printed_docx.render_docx(ps), worship_service.build_docx(**{**new_kwargs(), "liturgy": liturgy})):
        doc = Document(BytesIO(content))
        for text in ("1.\tone", "10.\tten", "100.\thundred", "999.\tmost"):
            fmt = _paragraph(doc, text).paragraph_format
            assert (fmt.left_indent, fmt.first_line_indent) == (Inches(0.375), -Inches(0.375)), text


def test_the_announcements_keep_a_heading_and_a_list_centered():
    week = sb.ServiceBulletin(announcements=sb.Announcements(other="# This week\n- Choir **at seven**"))
    word = Document(BytesIO(printed_docx.render_docx(service(bulletin=week))))
    heading, item = _paragraph(word, "This week"), _paragraph(word, "• Choir at seven")
    assert heading.alignment == item.alignment == printed_docx.WD_ALIGN_PARAGRAPH.CENTER
    assert [(r.bold, r.font.size) for r in heading.runs] == [(True, Pt(12))]


def test_the_printed_bulletin_s_box_headings_sit_under_its_element_headings():
    """The owner's answer 1 of 2026-10-25: bold 12 pt and 11 pt in the case typed, under the element's 11 pt bold
    capitals; the Word copy keeps 13 pt and 12 pt (test_the_word_copy_reads_the_marks)."""
    styles = printed_pdf.STYLES
    assert [styles[k].fontSize for k in ("element", "heading1", "heading2", "center_heading1", "center_heading2")] == [
        11, 12, 11, 12, 11]
    assert [printed_docx._SIZE[k] for k in ("heading1", "heading2", "center_heading1", "center_heading2")] == [
        12, 11, 12, 11]
    lines = pb.formatted("# For the world\n## Near")
    assert [(line.style, line.spans) for line in lines] == [("heading1", (Span("For the world", bold=True),)),
                                                             ("heading2", (Span("Near", bold=True),))]


@pytest.mark.parametrize("style", ["heading1", "heading2", "center_heading1", "center_heading2"])
def test_a_box_s_heading_at_the_foot_of_a_pdf_page_moves_on_with_what_follows(style):
    """As an element's heading does: a page break unless the heading and two lines fit (printed_pdf.KEEP)."""
    keep, _paragraph_ = printed_pdf._flowables(Line(style, (Span("Heading", bold=True),)), 300.0)
    assert isinstance(keep, CondPageBreak)
    own = printed_pdf.STYLES[style]
    assert keep.height >= own.spaceBefore + own.leading + 2 * printed_pdf.STYLES["body"].leading


def test_the_email_s_attachments_read_the_marks(church):
    files = email_usecase._attachments(church, _input(opening_prayer=FORMATTED), ["docx", "pdf"], None,
                                       lambda n: None)
    word = Document(BytesIO(files[0].content))
    assert [(r.bold, r.font.size) for r in _paragraph(word, "For the world").runs] == [(True, Pt(13))]
    assert "• the sick 12. the lonely" in _pdf_text(files[1].content)


def test_nothing_typed_becomes_markup(church):
    typed = "<b>x</b> & <font size=40>y</font> **<i>z</i> & <br/>**\n- <para>\n# </b>&amp;"
    pdf = _pdf_text(documents.build_printed(church, _input(opening_prayer=typed), "pdf", None,
                                            charge=lambda n: None).content)
    assert "<b>x</b> & <font size=40>y</font> <i>z</i> & <br/> • <para> </b>&amp;" in pdf
    for content in (documents.build_printed(church, _input(opening_prayer=typed), "docx", None,
                                            charge=lambda n: None).content,
                    documents.build_document(church, _input(opening_prayer=typed), "pastor").content):
        texts = [p.text for p in Document(BytesIO(content)).paragraphs]
        assert "<b>x</b> & <font size=40>y</font> <i>z</i> & <br/>" in texts and "</b>&amp;" in texts


def test_a_prayer_at_the_limit_full_of_marks_prints(church):
    text = ("- **a** *b*\n" * 2000)[: LIMITS.max_section_text]
    assert len(text) == LIMITS.max_section_text
    assert documents.build_printed(church, _input(prayers_of_the_people=text, opening_prayer=text), "pdf", None,
                                   charge=lambda n: None).content.startswith(b"%PDF")
    assert documents.build_document(church, _input(opening_prayer=text), "pastor").content.startswith(b"PK")


# --- the reviewer and the pews read the words, without the marks ---

def test_the_reviewer_and_the_pews_read_the_words_without_the_marks(tmp_db, make_user, make_church):
    owner = make_user(email="pastor@example.com")
    cid = make_church(name="Grace", owner_user_id=owner)
    client = make_api_client()
    cards = [{"section": "opening_prayer", "origin": "typed", "text": "# Gathering\nGracious God, as **we** journey, "
                                                                     "hear *us*.\n- all who mourn"},
             {"section": "call_to_worship", "origin": "typed", "text": "Leader: **Come**.\n# People: not a heading"}]
    r = client.post("/liturgy/review", json={"cards": cards}, headers=church_headers("pastor@example.com", cid))
    assert r.status_code == 200, r.text
    assert {"tag": "rules", "text": 'Stock phrase "as we journey". Say it more naturally.', "source": "code"} in \
        r.json()["cards"][0]["notes"]
    for path, reply in (("/liturgy/review", '{"cards": [], "service_notes": []}'), ("/liturgy/pews", '{"voices": []}')):
        fake = openai_client.FakeAI(reply=reply)
        openai_client.set_ai_for_tests(fake)
        r = client.post(path, json={"cards": cards}, headers=church_headers("pastor@example.com", cid))
        assert r.status_code == 200, r.text
        sent = fake.calls[0]["messages"][-1]["content"]
        assert "Gathering\nGracious God, as we journey, hear us.\n- all who mourn" in sent
        assert "Leader: Come.\n# People: not a heading" in sent
        assert "**" not in sent


def test_revise_gets_the_text_as_typed_marks_and_all(tmp_db, make_user, make_church):
    """S "AI interplay": Revise keeps the marks, so the AI is sent the draft with them."""
    cid = make_church(name="Grace", owner_user_id=make_user(email="pastor@example.com"))
    fake = openai_client.FakeAI(reply="Revised.")
    openai_client.set_ai_for_tests(fake)
    text = "# Gathering\nGracious God, hear **us** and *all* who mourn.\n- the sick\n2. the lonely"
    r = make_api_client().post("/liturgy/revise", json={"section": "opening_prayer", "text": text, "notes": ["Shorter."]},
                               headers=church_headers("pastor@example.com", cid))
    assert r.status_code == 200, r.text
    assert text in fake.calls[0]["messages"][-1]["content"]

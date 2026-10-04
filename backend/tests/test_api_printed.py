"""POST /documents/printed (printed bulletin spec, PR 1): the printed
bulletin as a PDF or a Word file, with the readings' text fetched in the
translation step 1 shows. scripture_fetcher.fetch_part is replaced, so
nothing leaves the machine."""
import logging
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

import scripture_fetcher
from api import ratelimit
from db import session_scope
from db.models import Church
from repos.memberships import add_membership
from scripture_fetcher import Part
from tests.api_helpers import (  # noqa: F401 (isolation_world is a fixture)
    assert_church_isolated,
    church_headers,
    isolation_world,
    make_api_client,
)

EMAIL = "pastor@example.com"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SERVICE = {
    "service_date_iso": "2026-10-04",
    "occasion": "World Communion Sunday",
    "scriptures": ["Isaiah 5:1-7", "Psalm 80:7-15", "Philippians 3:4b-14", "Matthew 21:33-46"],
    "hymns": {"opening": {"hymn_id": None, "title": "Old Favorite", "number": 12, "hymnal": "PH1990"},
              "response": None, "closing": None},
    "hymnal": None,
    "liturgy": {"call_to_worship": "Leader: Come. People: We come.", "prayers_of_the_people": "We pray."},
    "sermon_title": "Living Water",
    "selected_ot_ref": "",
    "selected_nt_ref": "",
    "include_communion": False,
    "custom_elements": [],
}


@pytest.fixture
def client(tmp_db):
    return make_api_client()


@pytest.fixture
def owner(make_user):
    return make_user(email=EMAIL)


@pytest.fixture
def church(owner, make_church):
    return make_church(name="Example Church", owner_user_id=owner)


@pytest.fixture
def calls(monkeypatch):
    """Every part fetched, as (part, translation); "Isaiah 5:1-7" is unavailable."""
    seen = []

    def fetch_part(part, translation):
        seen.append((part, translation))
        if part.startswith("Isaiah"):
            return Part(part, "unavailable", None)
        return Part(part, "ok", f"Text of {part} ({translation}).\nSecond verse.")

    monkeypatch.setattr(scripture_fetcher, "fetch_part", fetch_part)
    return seen


def post(client, church_id, body, email=EMAIL):
    return client.post("/documents/printed", json=body, headers=church_headers(email, church_id))


def pdf_text(content: bytes) -> str:
    return " ".join(" ".join(page.extract_text().split()) for page in PdfReader(BytesIO(content)).pages)


def test_the_pdf_downloads_with_the_file_headers_and_the_readings_text(client, church, calls, caplog):
    caplog.set_level(logging.INFO, logger="usecases.documents")
    r = post(client, church, {"format": "pdf", "translation": "kjv", "service": SERVICE})
    assert r.status_code == 200, r.text
    name = "printed_bulletin_October_04_2026.pdf"
    assert r.headers["content-type"] == PDF
    assert r.headers["content-disposition"] == f"attachment; filename=\"{name}\"; filename*=UTF-8''{name}"
    assert r.headers["cache-control"] == "no-store"
    text = pdf_text(r.content)
    assert "Example Church" in text and "*HYMN: #12 “Old Favorite”" in text
    assert "FIRST READING: Isaiah 5:1-7 [Reading text unavailable] NEW TESTAMENT" in text      # no names saved yet
    assert "NEW TESTAMENT READING: Philippians 3:4b-14 Text of Philippians 3:4-14 (kjv). Second verse." in text
    # Isaiah's text did not load, so only the New Testament reading is credited (build review fix M5).
    assert "Scripture readings are from" not in text
    assert "The New Testament Reading is from the King James Version (KJV)." in text
    assert calls == [("Isaiah 5:1-7", "kjv"), ("Philippians 3:4-14", "kjv")]
    assert "We pray." not in text
    (record,) = [r for r in caplog.records if r.getMessage().startswith("documents.printed")]
    assert "format=pdf" in record.getMessage() and "Living Water" not in record.getMessage()


def test_the_word_file_downloads_too(client, church, calls):
    r = post(client, church, {"format": "docx", "translation": None, "service": SERVICE})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX
    assert "printed_bulletin_October_04_2026.docx" in r.headers["content-disposition"]
    paragraphs = [p.text for p in Document(BytesIO(r.content)).paragraphs]
    assert "Text of Philippians 3:4-14 (web). Second verse." in paragraphs
    assert {translation for _part, translation in calls} == {"web"}


def test_a_translation_not_offered_here_prints_in_the_church_s(client, church, calls):
    with session_scope() as s:
        s.get(Church, church).settings = {"bible_translation": "asv"}
    for requested in ("esv", "nope", None):                  # no ESV key in tests
        calls.clear()
        r = post(client, church, {"format": "pdf", "translation": requested, "service": SERVICE})
        assert r.status_code == 200, r.text
        assert {translation for _part, translation in calls} == {"asv"}
        assert "American Standard Version (ASV)" in pdf_text(r.content)


def test_a_c1_control_character_in_a_free_text_does_not_print_as_a_question_mark(client, church, calls):
    """Build review fix M2: U+0085 is a line break, and the other C1 controls (outside cp1252, so the PDF
    printed them as "?") are deleted."""
    bulletin = {
        "prelude": {"title": "", "composer": ""}, "postlude": {"title": "", "composer": ""},
        "people": {"worship_leader": None, "liturgist": None, "organist": None}, "leaders": {},
        "announcements": {"ushers": "", "deacon": "", "coffee_hour": "", "activities": "Line one\x02\x85\x9bLine two",
                          "prayer_concerns": "", "collection": "", "other": ""},
        "reading_text": {"ot": "", "nt": "Pasted\x81 text."}, "unchecked": [],
    }
    r = post(client, church, {"format": "pdf", "service": {**SERVICE, "bulletin": bulletin}})
    assert r.status_code == 200, r.text
    text = pdf_text(r.content)
    assert "Line one Line two" in text and "Pasted text." in text
    assert "?" not in text


def test_no_readings_fetch_nothing(client, church, calls):
    r = post(client, church, {"format": "pdf", "service": {**SERVICE, "scriptures": []}})
    assert r.status_code == 200, r.text
    assert calls == []
    assert "Scripture readings are from" not in pdf_text(r.content)


def test_the_scripture_bucket_is_charged_per_part_and_a_429_fetches_nothing(client, church, owner, calls):
    ratelimit.consume("scripture", user_id=owner, cost=59)
    r = post(client, church, {"format": "pdf", "service": SERVICE})       # two parts, one token left
    assert r.status_code == 429, r.text
    assert r.json()["error"]["code"] == "rate_limited"
    assert calls == []


def test_any_member_and_only_members(client, church, make_user, calls, isolation_world):
    member = make_user(email="member@example.com")
    add_membership(member, church, "member")
    assert post(client, church, {"format": "docx", "service": SERVICE}, email="member@example.com").status_code == 200
    assert_church_isolated(client, "POST", "/documents/printed", world=isolation_world,
                           json={"format": "pdf", "service": SERVICE})


@pytest.mark.parametrize("change, field", [
    ({"format": "html"}, "format"),
    ({"translation": "x" * 21}, "translation"),
    ({"variant": "bulletin"}, "variant"),
    ({"service": {**SERVICE, "service_date_iso": "2026-02-30"}}, "service.service_date_iso"),
])
def test_a_bad_body_is_a_422_naming_the_field(client, church, calls, change, field):
    r = post(client, church, {"format": "pdf", "service": SERVICE, **change})
    assert r.status_code == 422, r.text
    assert field in r.json()["error"]["fields"], r.text
    assert calls == []


def test_the_church_s_bulletin_settings_print_for_every_member(client, church, make_user, calls):
    """PR 2a: the details an admin saved print on everyone's bulletin; a blank one prints nothing."""
    with session_scope() as s:
        s.get(Church, church).settings = {"bulletin": {
            "address_lines": ["100 Example Street"], "phone": "", "facebook": "Example Church",
            "service_time": "10:30 a.m.", "worship_leader": "Rev. Alex Example", "liturgist": "Sam Sample",
            "organist": "", "starred": ["sermon"], "leaders": {"sermon": "worship_leader", "nt_reading": "liturgist"},
            "stand_note": "Please stand if able", "gloria_patri_words": "Glory be."}}
    member = make_user(email="member@example.com")
    add_membership(member, church, "member")
    r = post(client, church, {"format": "pdf", "service": SERVICE}, email="member@example.com")
    assert r.status_code == 200, r.text
    text = pdf_text(r.content)
    assert "October 4, 2026 100 Example Street FB: Example Church 1 THE SERVICE FOR THE LORD’S DAY Example Church " \
           "Rev. Alex Example, Worship Leader Sam Sample, Liturgist October 4, 2026 10:30 a.m." in text
    assert "NEW TESTAMENT READING: Philippians 3:4b-14 Sam Sample" in text
    assert "*SERMON: “Living Water” Rev. Alex Example" in text and "*HYMN" not in text
    assert "Glory be." in text and text.count("*Please stand if able") == 1
    for gone in ("Organist", "[Organist]", "[Liturgist]", "[Worship leader]", "[Service time]", "[Phone]", "[Email]"):
        assert gone not in text, gone
    # No bulletin was posted (a page from before the Bulletin step): PR 1's weekly placeholders (PR 2b-1).
    assert "‘[Prelude title]’" in text and "Coffee Hour: [Name]" in text


def test_a_stored_control_character_still_prints_in_word(client, church, calls):
    """PR 2a: settings written by another path are made Word-safe when printed, as the PUT makes them."""
    with session_scope() as s:
        s.get(Church, church).settings = {"bulletin": {"organist": "Jordan\x01 Doe", "phone": "(555)\x0b010-0100"}}
    r = post(client, church, {"format": "docx", "service": SERVICE})
    assert r.status_code == 200, r.text
    paragraphs = [p.text for p in Document(BytesIO(r.content)).paragraphs]
    assert "Jordan Doe, Organist" in paragraphs and "(555) 010-0100" in paragraphs   # one line (build review I1)


def test_the_week_s_fields_print_and_a_pasted_reading_is_neither_fetched_nor_charged(client, church, owner,
                                                                                  calls, caplog):
    """PR 2b: the music, this week's people and part leaders, the announcements and a pasted reading."""
    with session_scope() as s:
        s.get(Church, church).settings = {"bulletin": {"liturgist": "Sam Sample", "organist": "Jordan Doe"}}
    bulletin = {
        "prelude": {"title": "Morning Voluntary", "composer": "Pat Example"},
        "postlude": {"title": "Festive Postlude", "composer": ""},
        "people": {"worship_leader": "Rev. Guest", "liturgist": "", "organist": None},
        "leaders": {"sermon": "Pat Example"},
        "announcements": {"ushers": "Sam Sample", "deacon": "", "coffee_hour": "The Example family",
                          "activities": "", "prayer_concerns": "For all who are ill.", "collection": "",
                          "other": ""},
        "reading_text": {"ot": "", "nt": "Pasted text of the reading."},
        "unchecked": ["prelude"],
    }
    ratelimit.consume("scripture", user_id=owner, cost=59)                   # one token left: one part
    caplog.set_level(logging.INFO, logger="usecases.documents")
    r = post(client, church, {"format": "pdf", "service": {**SERVICE, "bulletin": bulletin}})
    assert r.status_code == 200, r.text
    assert calls == [("Isaiah 5:1-7", "web")]                                # the pasted one is not fetched
    text = pdf_text(r.content)
    for expected in ("Rev. Guest, Worship Leader Jordan Doe, Organist October 4, 2026",
                     "PRELUDE: ‘Morning Voluntary’ Jordan Doe - Pat Example", "SERMON: “Living Water” Pat Example",
                     "NEW TESTAMENT READING: Philippians 3:4b-14 Rev. Guest Pasted text of the reading.",
                     "POSTLUDE: ‘Festive Postlude’ Jordan Doe", "ANNOUNCEMENTS October 4, 2026",
                     "Ushers/Counters: Sam Sample Coffee Hour: The Example family PRAYERS AND CONCERNS "
                     "For all who are ill."):
        assert expected in text, expected
    assert "Sam Sample, Liturgist" not in text and "Scripture readings are from" not in text
    # The only fetched reading (Isaiah) did not load: no credit line at all (build review fix M5).
    assert " is from the " not in text
    (record,) = [r for r in caplog.records if r.getMessage().startswith("documents.printed")]
    assert "ill" not in record.getMessage() and "Example" not in record.getMessage()



def test_the_cover_picture_prints_and_a_page_from_before_3b_keeps_pr_1_s_box(client, church, owner, make_church,
                                                                             calls):
    """PR 3a: the church's picture fills the cover's box (the reading and the date on its band); null (no
    picture this week) or a picture the church does not have prints the reading and the date alone; a
    bulletin that does not say (a page from before PR 3b) keeps PR 1's [Cover picture] box."""
    from tests.picture_helpers import picture

    def picture_id(church_id):
        r = client.post("/bulletin-images", content=picture((1600, 1200)),
                        headers={**church_headers(EMAIL, church_id), "Content-Type": "image/jpeg"})
        assert r.status_code == 201, r.text
        return r.json()["id"]

    mine, theirs = picture_id(church), picture_id(make_church(name="Other", owner_user_id=owner))
    blank = {"prelude": {"title": "", "composer": ""}, "postlude": {"title": "", "composer": ""},
             "people": {"worship_leader": None, "liturgist": None, "organist": None}, "leaders": {},
             "announcements": {key: "" for key in ("ushers", "deacon", "coffee_hour", "activities",
                                                   "prayer_concerns", "collection", "other")},
             "reading_text": {"ot": "", "nt": ""}, "unchecked": []}
    for cover, images, starts in ((mine, 1, "Example Church 1 THE SERVICE"),
                                  (None, 0, "Example Church Philippians 3:4b-14 October 4, 2026 1 THE SERVICE"),
                                  (theirs, 0, "Example Church Philippians 3:4b-14 October 4, 2026 1 THE SERVICE"),
                                  ("left out", 0, "Example Church [Cover picture] Philippians 3:4b-14")):
        bulletin = blank if cover == "left out" else {**blank, "cover_image_id": cover}
        r = post(client, church, {"format": "pdf", "service": {**SERVICE, "bulletin": bulletin}})
        assert r.status_code == 200, r.text
        first = PdfReader(BytesIO(r.content)).pages[0]
        assert len(first.images) == images, cover
        assert pdf_text(r.content).startswith(starts), cover
    r = post(client, church, {"format": "docx", "service": {**SERVICE, "bulletin": {**blank, "cover_image_id": mine}}})
    assert r.status_code == 200, r.text
    assert len(Document(BytesIO(r.content)).inline_shapes) == 1

"""The printed bulletin's two files (printed bulletin spec, PR 1): the
print-ready PDF (printed_pdf) and the Word file (printed_docx), read back
with pypdf and python-docx. PR 2b: the week's music and announcements, and
no announcements page when every announcement is blank."""
import dataclasses
import datetime
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader

import bulletin_image
import bulletin_settings as bs
import printed_bulletin as pb
import printed_docx
import printed_pdf
from service_output import ResolvedHymn, ResolvedService
import service_bulletin as sb
from tests import picture_helpers
from tests.test_printed_bulletin import SETTINGS, WEEK

VERSE = "And he answered, I will not; but afterward he repented and went. "


def service(verses: int = 20) -> pb.PrintedService:
    resolved = ResolvedService(
        service_date=datetime.date(2026, 9, 27), scriptures=("Psalm 25:1-9", "Matthew 21:23-32"),
        hymns={"opening": ResolvedHymn("God Is Here!", 409), "response": None,
               "closing": ResolvedHymn("Jesus Shall Reign Where’er the Sun", 265)},
        liturgy={"call_to_worship": "Leader: Lift up your hearts. People: We come, ready to listen and learn.",
                 "opening_prayer": "God of wisdom and truth, hear us. Amen",
                 "benediction": "Go in peace. - A Friend"},
        sermon_title="Who Said?")
    return pb.PrintedService("Example Church", resolved, pb.Reading("Psalm 25:1-9", "In you, Lord, I put my trust."),
                             pb.Reading("Matthew 21:23-32", VERSE * verses), "World English Bible (WEB)", SETTINGS,
                             WEEK)


def sides(content: bytes) -> list[str]:
    return [" ".join(page.extract_text().split()) for page in PdfReader(BytesIO(content)).pages]


def halves(content: bytes) -> list[str]:
    """Each side's left and right booklet page, by where the text is drawn."""
    out = []
    for page in PdfReader(BytesIO(content)).pages:
        parts: list[list[str]] = [[], []]

        def visit(text, cm, tm, _font, _size):
            if text.strip():
                parts[tm[4] * cm[0] + tm[5] * cm[2] + cm[4] >= pb.PAGE_WIDTH].append(text)

        page.extract_text(visitor_text=visit)
        out += [" ".join(" ".join(half).split()) for half in parts]
    return out


def test_the_pdf_is_legal_landscape_sides_with_two_pages_each_in_reading_order():
    content = printed_pdf.render_pdf(service())
    assert content.startswith(b"%PDF-")
    reader = PdfReader(BytesIO(content))
    assert [(float(p.mediabox.width), float(p.mediabox.height)) for p in reader.pages] == [(1008.0, 612.0)] * 2
    pages = halves(content)
    # Side 1: the cover (no number), then page 1; side 2: pages 2 and 3 (the announcements, last).
    assert pages[0].startswith("Example Church [Cover picture] Matthew 21:23-32 September 27, 2026 100 Example Street")
    assert pages[0].endswith("FB: Example Church")
    assert pages[1].startswith("1 THE SERVICE FOR THE LORD’S DAY Example Church Rev. Alex Example, Worship Leader")
    # The long reading starts under its heading on page 1 and runs on to page 2 (no gap before it).
    assert f"NEW TESTAMENT READING: Matthew 21:23-32 Rev. Alex Example {VERSE.strip()}" in pages[1]
    assert pages[2].startswith("2 And he answered") and pages[2].endswith("*Congregation stands if able")
    assert pages[3].startswith("3 ANNOUNCEMENTS September 27, 2026 Ushers/Counters: Sam Sample, Jordan Doe")
    assert pages[3].endswith("OTHER ANNOUNCEMENTS The office is closed on Monday.")


def test_the_pdf_prints_the_service_and_its_readings():
    text = " ".join(sides(printed_pdf.render_pdf(service())))
    for expected in ("*HYMN: #409 “God Is Here!”", "*HYMN: #265 “Jesus Shall Reign Where’er the Sun”",
                     "CALL TO WORSHIP Sam Sample", "Leader: Lift up your hearts.",
                     "People: We come, ready to listen and learn.", "FIRST READING: Psalm 25:1-9",
                     "In you, Lord, I put my trust.", "NEW TESTAMENT READING: Matthew 21:23-32 Rev. Alex Example",
                     "Scripture readings are from the World English Bible (WEB).", "SERMON: “Who Said?”",
                     "I believe in God, the Father almighty", "*Congregation stands if able",
                     "POSTLUDE: ‘Festive Postlude’ Jordan Doe", "- Lee Sample", "Go in peace. - A Friend",
                     "Tuesday: Bible study at 10 a.m. Wednesday: Choir at 7 p.m."):
        assert expected in text, expected


def test_any_page_count_takes_half_as_many_sides_rounded_up_with_the_announcements_last():
    for verses, count in ((100, 6), (300, 9)):               # an even and an odd page count
        pages = halves(printed_pdf.render_pdf(service(verses)))
        assert len(pages) == 2 * -(-count // 2), verses       # ceil(count / 2) sides, two halves each
        printed, blank = pages[:count], pages[count:]
        assert [page.split(" ", 1)[0] for page in printed[1:]] == [str(n) for n in range(1, count)]
        assert printed[-1].startswith(f"{count - 1} ANNOUNCEMENTS") and blank == [""] * (count % 2)


def test_characters_the_standard_fonts_cannot_print_become_a_question_mark():
    assert printed_pdf.to_pdf_text("“Grace” – ﬁne ש") == "“Grace” – fine ?"
    assert printed_pdf.to_pdf_text("½ cup… é a\u200bb\ufeff") == "½ cup… é ab"     # pasted invisibles drop
    resolved = ResolvedService(service_date=datetime.date(2026, 9, 27), sermon_title="Shalom שלום")
    text = " ".join(sides(printed_pdf.render_pdf(pb.PrintedService("Example Church", resolved))))
    assert "SERMON: “Shalom ????”" in text


def test_the_word_file_is_the_same_booklet_in_reading_order():
    doc = Document(BytesIO(printed_docx.render_docx(service())))
    section = doc.sections[0]
    assert (section.page_width, section.page_height) == (Inches(7), Inches(8.5))
    assert section.left_margin == Inches(0.5) and section.different_first_page_header_footer
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs[0] == "Example Church"
    assert doc.tables[0].cell(0, 0).paragraphs[1].text == "Matthew 21:23-32"
    starts = [p.text for p in doc.paragraphs if p.paragraph_format.page_break_before]
    assert starts == ["THE SERVICE FOR THE LORD’S DAY", "ANNOUNCEMENTS"]
    for expected in ("*HYMN:  #409  “God Is Here!”", "CALL TO WORSHIP\tSam Sample",
                     "People: We come, ready to listen and learn.", "FIRST READING:  Psalm 25:1-9\tSam Sample",
                     "Scripture readings are from the World English Bible (WEB).", "*Congregation stands if able",
                     "FB: Example Church", "Coffee Hour: The Example family",
                     "Tuesday: Bible study at 10 a.m.\nWednesday: Choir at 7 p.m."):
        assert expected in paragraphs, expected
    people = next(p for p in doc.paragraphs if p.text.startswith("People: We come"))
    assert all(run.bold for run in people.runs)
    assert 'w:instrText xml:space="preserve">PAGE<' in section.footer._element.xml
    # The page numbers start at 0 on the cover, so the first inside page is 1; pgNumType sits in the
    # schema's order (after pgMar, before cols and titlePg), not at the end (PR 1 build review fix 1).
    children = [child.tag.split("}")[1] for child in section._sectPr]
    assert children == ["footerReference", "pgSz", "pgMar", "pgNumType", "cols", "titlePg", "docGrid"]
    assert section._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "0"


def test_the_longest_details_still_fit_the_cover():
    """PR 2a: at every limit, with a church name of two lines, the contact lines shrink to fit the cover, so
    the order of worship still starts on page 1 (side 1 stays the cover and page 1)."""
    def longest(field: str, word: str) -> str:
        return (word * bs.MAX_LENGTH[field])[:bs.MAX_LENGTH[field]].strip()

    settings = bs.BulletinSettings(
        address_lines=tuple(longest("address_line", f"{n}00 Example Street ") for n in (1, 2, 3)),
        phone=longest("phone", "(555) 010-0100 "), email=longest("email", "office.")[:-12] + "@example.com",
        website=longest("website", "example.com/"), facebook=longest("facebook", "Example Church "),
        **{role: longest("person", "Alex Example ") for role in bs.ROLES})
    ps = dataclasses.replace(service(), church_name="The First Presbyterian Church of Springfield", settings=settings)
    pages = halves(printed_pdf.render_pdf(ps))
    assert pages[0].startswith("The First Presbyterian Church of Springfield [Cover picture]")
    assert pages[0].endswith(f"FB: {settings.facebook}")
    assert pages[1].startswith("1 THE SERVICE FOR THE LORD’S DAY The First Presbyterian Church of Springfield")



@pytest.mark.parametrize("length, contact, cover_ends", [
    (145, True, "FB: Example Church"),
    (150, False, "[Cover picture] Matthew 21:23-32 September 27, 2026"),
    (150, True, "[Cover picture] Matthew 21:23-32 September 27, 2026"),
])
def test_a_long_church_name_still_leaves_page_1_on_side_1(length, contact, cover_ends):
    """Build review M1: under a church name of 5 title lines (about 145 characters) the contact lines shrink
    into the room left. A name of 6 lines (150; names may have 200) leaves no room under the picture: nothing
    is kept for contact lines (none, or no room for them), so the order of worship still starts on side 1."""
    name = ("Saint Example " * 12)[:length].strip()
    settings = SETTINGS if contact else bs.BulletinSettings()
    content = printed_pdf.render_pdf(dataclasses.replace(service(), church_name=name, settings=settings))
    assert len(PdfReader(BytesIO(content)).pages) == 2
    pages = halves(content)
    assert pages[0].startswith("Saint Example") and pages[0].endswith(cover_ends)
    assert pages[1].startswith("1 THE SERVICE FOR THE LORD’S DAY")


def test_the_word_cover_picture_box_is_centered():
    """Owner's desktop Word check (2026-10-03): the picture box sat at the left margin, off center under the
    centered church name. The table is centered (w:jc in tblPr, in the schema's order)."""
    box = Document(BytesIO(printed_docx.render_docx(service()))).tables[0]
    tbl_pr = box._tbl.tblPr
    assert tbl_pr.find(qn("w:jc")).get(qn("w:val")) == "center"
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
             "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    children = [child.tag.split("}")[1] for child in tbl_pr]
    assert children == sorted(children, key=order.index)


def test_with_every_announcement_blank_the_announcements_page_is_left_out():
    """PR 2b: a blank announcement prints nothing, and with none at all there is no announcements page; the
    order of worship's last page is the last page (an odd count leaves the last right half blank)."""
    ps = dataclasses.replace(service(), bulletin=sb.ServiceBulletin(prelude=WEEK.prelude))
    pages = halves(printed_pdf.render_pdf(ps))
    assert len(pages) == 4 and pages[3] == ""
    assert pages[2].startswith("2 And he answered") and pages[2].endswith("*Congregation stands if able")
    assert "ANNOUNCEMENTS September" not in " ".join(pages)
    doc = Document(BytesIO(printed_docx.render_docx(ps)))
    assert [p.text for p in doc.paragraphs if p.paragraph_format.page_break_before] == [
        "THE SERVICE FOR THE LORD’S DAY"]
    assert doc.paragraphs[-1].text == "*Congregation stands if able"


def chosen(picture: bytes | None = None, **cover) -> pb.PrintedService:
    week = dataclasses.replace(WEEK, cover_given=True, **cover)
    return dataclasses.replace(service(), bulletin=week, cover_picture=picture)


def test_the_cover_picture_fills_the_box_in_both_files():
    """PR 3 planning answers 2-4: the same picture in the PDF and the Word version, 372 x 300 pt under the
    church's name, the reading and the date on its band (inside the picture, so not text)."""
    ps = chosen(bulletin_image.prepare(picture_helpers.stripes((1600, 1200), [picture_helpers.GREEN])).content,
                cover_image_id="0b4c2b0e-1111-4222-8333-444455556666")
    content = printed_pdf.render_pdf(ps)
    (image,) = PdfReader(BytesIO(content)).pages[0].images
    assert image.image.size == (1550, 1250)
    assert halves(content)[0].startswith("Example Church 100 Example Street")
    doc = Document(BytesIO(printed_docx.render_docx(ps)))
    (shape,) = doc.inline_shapes
    assert (shape.width, shape.height) == (Pt(372), Pt(300))
    assert doc.tables == [] and "[Cover picture]" not in [p.text for p in doc.paragraphs]
    word_picture = doc.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob
    assert word_picture == printed_pdf.cover_picture(ps)                   # the PDF's picture, byte for byte


def test_no_picture_prints_the_reading_and_the_date_where_it_would_be():
    """PR 3 planning answer 6: no box at all; the reading and the date centered in its place."""
    ps = chosen()
    content = printed_pdf.render_pdf(ps)
    assert len(PdfReader(BytesIO(content)).pages[0].images) == 0
    assert halves(content)[0].startswith("Example Church Matthew 21:23-32 September 27, 2026 100 Example Street")
    doc = Document(BytesIO(printed_docx.render_docx(ps)))
    assert doc.tables == [] and len(doc.inline_shapes) == 0
    reading, date = doc.paragraphs[1:3]
    assert (reading.text, date.text) == ("Matthew 21:23-32", "September 27, 2026")
    assert reading.runs[0].bold and reading.runs[0].font.size == Pt(18) and date.runs[0].font.size == Pt(14)
    assert reading.paragraph_format.space_before == date.paragraph_format.space_after > Pt(100)

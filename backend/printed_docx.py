"""The printed bulletin as an editable Word file (printed bulletin spec, PR 1).

The same pages as printed_pdf in reading order, one 7 x 8.5 in page each
(half a legal sheet), 0.5 in margins, Times New Roman 11 pt: the cover, the
order of worship and the announcements (left out when every announcement is
blank, PR 2b), each starting a page, numbered from the first inside page.
A free text's line breaks stay line breaks (python-docx turns "\n" into
one). The cover's box (PR 3a) is the PDF's: the same picture
(printed_pdf.cover_picture), the reading and the date alone, or PR 1's
[Cover picture] box. An element's leader sits at a right tab stop. To print
it as the PDF prints, two pages to a legal sheet, use the printer's "2 pages
per sheet" setting; the PDF is already arranged that way.
Pure: no database or FastAPI.
"""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import printed_bulletin as pb
import printed_pdf

FONT = "Times New Roman"
TEXT_WIDTH = Inches(6)                   # 7 in less two 0.5 in margins

_ALIGN = {"header": WD_ALIGN_PARAGRAPH.CENTER, "section": WD_ALIGN_PARAGRAPH.CENTER,
          "title": WD_ALIGN_PARAGRAPH.CENTER, "box": WD_ALIGN_PARAGRAPH.CENTER,
          "contact": WD_ALIGN_PARAGRAPH.CENTER, "center": WD_ALIGN_PARAGRAPH.CENTER}
_SPACE_BEFORE = {"element": 8, "section": 12, "note": 10, "credit": 4, "center": 3}
_SIZE = {"title": 28, "credit": 9, "contact": 12}
# The w:sectPr children that follow w:pgNumType (ECMA-376 CT_SectPr).
_AFTER_PG_NUM_TYPE = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi",
                      "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")


def _add_line(doc, line: pb.Line):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_after = Pt(4 if line.style == "section" else 0)
    fmt.space_before = Pt(_SPACE_BEFORE.get(line.style, 0))
    if line.style in _ALIGN:
        p.alignment = _ALIGN[line.style]
    if line.style == "hanging":
        fmt.left_indent, fmt.first_line_indent = Inches(0.5), Inches(-0.5)
    elif line.style == "indent":
        fmt.left_indent = Inches(0.5)
    if line.style == "element":
        fmt.keep_with_next = True
    for span in line.spans:
        run = p.add_run(span.text)
        run.bold, run.italic = span.bold or None, span.italic or None
        if line.style in _SIZE:
            run.font.size = Pt(_SIZE[line.style])
        if line.style == "contact":
            run.font.name = "Arial"
    if line.right:
        fmt.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t" + line.right)
    return p


def _page_number_footer(section) -> None:
    """A centered PAGE field; the cover (a different first page) shows none,
    and numbering starts at 0 there, so the first inside page is 1."""
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.bold = True
    for tag, text in (("w:fldChar", "begin"), ("w:instrText", "PAGE"), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), text)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    start = OxmlElement("w:pgNumType")
    start.set(qn("w:start"), "0")
    # In the schema's place (before cols, titlePg, docGrid), not appended at the end: Word reads sectPr
    # strictly in order (PR 1 build review fix 1).
    section._sectPr.insert_element_before(start, *_AFTER_PG_NUM_TYPE)


def _cover_box(doc, ps: pb.PrintedService, label: pb.Line, reference: pb.Line, date: pb.Line) -> None:
    """The cover's box for its kind (printed_bulletin.cover_kind), as the PDF's."""
    kind = pb.cover_kind(ps)
    width, height = pb.COVER_BOX
    if kind == "picture":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(BytesIO(printed_pdf.cover_picture(ps)), width=Pt(width), height=Pt(height))
        return
    if kind == "none":                     # no box: the reading and the date centered where the picture would be
        lines = [(line.text, size) for line, size in zip((reference, date), pb.COVER_TEXT_POINTS) if line.text]
        paragraphs = []
        for text, size in lines:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold, run.font.size = True, Pt(size)
            paragraphs.append(p)
        room = Pt((height - sum(size * 1.2 for _text, size in lines)) / 2)
        paragraphs[0].paragraph_format.space_before = paragraphs[-1].paragraph_format.space_after = room
        return
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.alignment = WD_TABLE_ALIGNMENT.CENTER      # centered under the name (owner's desktop Word check, 2026-10-03)
    cell = box.cell(0, 0)
    cell.width = Inches(5.2)
    first = cell.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hint = first.add_run(label.text)
    hint.italic = True
    hint.font.color.rgb = RGBColor(0x73, 0x73, 0x73)
    for _ in range(8):
        first.add_run().add_break()
    for line in (reference, date):
        p = cell.add_paragraph(line.text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.add_paragraph()


def render_docx(ps: pb.PrintedService) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(7), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.5))
    section.footer_distance = Inches(0.25)
    _page_number_footer(section)

    title, label, reference, date, *contact = pb.cover(ps)
    _add_line(doc, title)
    _cover_box(doc, ps, label, reference, date)
    doc.add_paragraph()
    for line in contact:
        _add_line(doc, line)

    for part in (pb.order_of_worship(ps), pb.announcements(ps)):
        if not part:                                   # every announcement blank: no page (PR 2b)
            continue
        first, *rest = part
        _add_line(doc, first).paragraph_format.page_break_before = True
        for line in rest:
            _add_line(doc, line)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()

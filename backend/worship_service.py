#!/usr/bin/env python3
"""
Worship service generator: the Word document export. The hymn helpers moved
to hymn_search, hymn_suggest and usecases.hymns in slice 3; liturgy generation
moved to liturgy_prompts and usecases.liturgy in slice 4. Slice 5a renders
the files through service_output.render_docx (POST /documents).
"""

from __future__ import annotations

import logging
import re
from collections.abc import Mapping, Sequence
from typing import Dict, Any, List, Optional
from io import BytesIO

from liturgy_config import ASSURANCE_RESPONSE, COMMUNION_BLOCKS
from service_output import hymn_line

logger = logging.getLogger(__name__)

# Optional import for docx
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None


def _add_leader_people_paragraph(doc, text: str) -> None:
    """Add paragraph(s): each Leader: (normal), each People: (bold). Supports multiple Leader/People pairs."""
    if not text or not text.strip():
        return
    # Find all "Leader:" and "People:" in order (case-insensitive)
    pattern = re.compile(r"\b(Leader|People):\s*", re.IGNORECASE)
    pos = 0
    parts = []
    for m in pattern.finditer(text):
        if m.start() > pos:
            parts.append(("", text[pos : m.start()].strip()))  # preamble if any
        role = "Leader" if m.group(1).lower() == "leader" else "People"
        end = pattern.search(text, m.end())
        content_end = end.start() if end else len(text)
        content = text[m.end() : content_end].strip()
        parts.append((role, content))
        pos = content_end
    if not parts:
        doc.add_paragraph(text)
        return
    for role, content in parts:
        if not content and role == "":
            continue
        if role == "People":
            p = doc.add_paragraph()
            p.add_run("People: ")
            r = p.add_run(content)
            r.bold = True
        else:
            line = ("Leader: " + content) if role == "Leader" else content
            if line:
                doc.add_paragraph(line)


def _add_communion_liturgy(doc) -> None:
    """Add The Sacrament of the Lord's Supper liturgy: liturgy_config.COMMUNION_BLOCKS
    (the one copy of the text, which GET /liturgy/config also serves), one paragraph each."""
    for block in COMMUNION_BLOCKS:
        if block.style == "heading1":
            doc.add_paragraph(block.text, style="Heading 1")
        elif block.style == "heading2":
            doc.add_paragraph(block.text, style="Heading 2")
        elif block.style == "response":
            doc.add_paragraph().add_run(block.text).bold = True
        elif block.style == "text":
            doc.add_paragraph(block.text)
        else:                               # blank
            doc.add_paragraph()


def _add_assurance_paragraph(doc, leader_text: str) -> None:
    """Add Assurance: Leader line then liturgy_config.ASSURANCE_RESPONSE in bold
    (the one copy the 4b card shows too)."""
    leader_clean = (leader_text or "").strip()
    if leader_clean.startswith("Leader:"):
        leader_clean = leader_clean[7:].strip()
    if leader_clean:
        doc.add_paragraph("Leader: " + leader_clean)
    # Always add the congregational response
    p = doc.add_paragraph()
    r = p.add_run(ASSURANCE_RESPONSE)
    r.bold = True


def _add_custom_elements_after(
    doc,
    anchor: str,
    custom_elements: List[Dict[str, Any]],
) -> None:
    """Add any custom elements that are inserted after this anchor."""
    for ce in custom_elements:
        if ce.get("insert_after") == anchor and ce.get("label"):
            doc.add_paragraph(ce["label"], style="Heading 2")
            if ce.get("text"):
                doc.add_paragraph(ce["text"])
            doc.add_paragraph()


def _add_hymn(doc, heading: str, hymn: Optional[Mapping[str, Any]]) -> None:
    """A filled slot's heading and line; an empty slot prints nothing (5a: headings by slot)."""
    if hymn is None:
        return
    doc.add_paragraph(heading, style="Heading 2")
    doc.add_paragraph(hymn_line(hymn.get("title") or "", hymn.get("number")))
    doc.add_paragraph()


def build_docx(
    *,
    occasion: str,
    date_display: str,
    hymns_by_slot: Mapping[str, Optional[Mapping[str, Any]]],
    liturgy: Mapping[str, str],
    ot_ref: Optional[str],
    nt_ref: Optional[str],
    sermon_title: str = "",
    include_sermon: bool = True,
    include_prayers_of_the_people: bool = True,
    include_communion: bool = False,
    custom_elements: Sequence[Mapping[str, Any]] = (),
) -> bytes:
    """
    The Word file of a service: Streamlit's layout, paragraph for paragraph
    (Times New Roman 11 pt, Word's Heading 2, the bold People lines and
    confession), with slice 5a's changes: hymn headings follow slots
    (opening "First Hymn", response "Second Hymn", closing "Third Hymn"; an
    empty slot prints nothing), a hymn without a number prints no "#None", and
    the first reading's heading is "First Reading". ot_ref and nt_ref arrive
    resolved (service_output.resolve_doc_readings); None leaves a reading out.
    Both variants include the sermon title; only the pastor's copy includes
    Prayers of the People (service_output.VARIANTS). Custom elements print
    after their anchor, which is emitted even when its item is absent.
    Returns the .docx bytes.
    """
    if not Document:
        raise RuntimeError("python-docx is required. pip install python-docx")

    custom = list(custom_elements)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Times New Roman"

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Worship Service\n{occasion}")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"
    if date_display:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(date_display).font.size = Pt(12)
    doc.add_paragraph()

    # 1. Call to Worship (Leader / People, People in bold)
    if liturgy.get("call_to_worship"):
        doc.add_paragraph("Call to Worship", style="Heading 2")
        _add_leader_people_paragraph(doc, liturgy["call_to_worship"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "call_to_worship", custom)

    # 2. Opening Prayer
    if liturgy.get("opening_prayer"):
        doc.add_paragraph("Opening Prayer", style="Heading 2")
        doc.add_paragraph(liturgy["opening_prayer"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "opening_prayer", custom)

    # 3. First Hymn (the opening slot)
    _add_hymn(doc, "First Hymn", hymns_by_slot.get("opening"))
    _add_custom_elements_after(doc, "first_hymn", custom)

    # 4. Prayer of Confession (bold)
    if liturgy.get("prayer_of_confession"):
        doc.add_paragraph("Prayer of Confession", style="Heading 2")
        p = doc.add_paragraph()
        p.add_run(liturgy["prayer_of_confession"]).bold = True
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayer_of_confession", custom)

    # 5. Assurance of Pardon (Leader: ... / ASSURANCE_RESPONSE in bold)
    if liturgy.get("assurance"):
        doc.add_paragraph("Assurance of Pardon", style="Heading 2")
        _add_assurance_paragraph(doc, liturgy["assurance"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "assurance", custom)

    # 6. Prayer for Illumination
    if liturgy.get("prayer_for_illumination"):
        doc.add_paragraph("Prayer for Illumination", style="Heading 2")
        doc.add_paragraph(liturgy["prayer_for_illumination"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayer_for_illumination", custom)

    # 7. First Reading (reference only; owner decision B renamed "Old Testament Reading")
    if ot_ref:
        doc.add_paragraph("First Reading", style="Heading 2")
        doc.add_paragraph(ot_ref)
        doc.add_paragraph()
    _add_custom_elements_after(doc, "ot_reading", custom)

    # 8. New Testament Reading (reference only)
    if nt_ref:
        doc.add_paragraph("New Testament Reading", style="Heading 2")
        doc.add_paragraph(nt_ref)
        doc.add_paragraph()
    _add_custom_elements_after(doc, "nt_reading", custom)

    # 9. Sermon Title (both copies)
    if include_sermon:
        doc.add_paragraph("Sermon Title", style="Heading 2")
        doc.add_paragraph(sermon_title.strip() if sermon_title and sermon_title.strip() else "[Sermon title]")
        doc.add_paragraph()
    _add_custom_elements_after(doc, "sermon", custom)

    # 10. Affirmation of Faith
    doc.add_paragraph("Affirmation of Faith", style="Heading 2")
    doc.add_paragraph("Apostles' Creed")
    doc.add_paragraph()
    _add_custom_elements_after(doc, "affirmation_of_faith", custom)

    # 11. Second Hymn (the response slot)
    _add_hymn(doc, "Second Hymn", hymns_by_slot.get("response"))
    _add_custom_elements_after(doc, "second_hymn", custom)

    # 11b. Communion liturgy (after the second hymn when communion is included)
    if include_communion:
        _add_communion_liturgy(doc)
    _add_custom_elements_after(doc, "communion", custom)

    # 12. Prayers of the People (the pastor's copy only)
    if include_prayers_of_the_people and liturgy.get("prayers_of_the_people"):
        doc.add_paragraph("Prayers of the People", style="Heading 2")
        doc.add_paragraph(liturgy["prayers_of_the_people"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "prayers_of_the_people", custom)

    # 13. Offertory Prayer
    if liturgy.get("offertory_prayer"):
        doc.add_paragraph("Offertory Prayer", style="Heading 2")
        doc.add_paragraph(liturgy["offertory_prayer"])
        doc.add_paragraph()
    _add_custom_elements_after(doc, "offertory_prayer", custom)

    # 14. Third Hymn (the closing slot)
    _add_hymn(doc, "Third Hymn", hymns_by_slot.get("closing"))
    _add_custom_elements_after(doc, "third_hymn", custom)
    _add_custom_elements_after(doc, "benediction", custom)  # "Before Benediction"

    # Benediction
    if liturgy.get("benediction"):
        doc.add_paragraph("Benediction", style="Heading 2")
        doc.add_paragraph(liturgy["benediction"])

    _add_custom_elements_after(doc, "end", custom)  # At the end (after Benediction)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
